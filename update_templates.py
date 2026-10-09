import openpyxl
import pandas as pd
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

print("Loading DB_HocSinh to create Name -> ID mapping...")
df_hs = pd.read_excel('DB_HocSinh.xlsx')
# Create a mapping. If there are duplicates, we might need to handle them carefully.
# For simplicity, we just use the first match.
name_to_id = {}
for _, row in df_hs.iterrows():
    name = str(row['HoTen']).strip().lower()
    ma_hs = str(row['MaHS']).strip()
    if name not in name_to_id:
        name_to_id[name] = ma_hs

def update_excel_file(file_path, output_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return
        
    print(f"Processing {file_path}...")
    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        print(f"Failed to open {file_path}: {e}")
        return

    # To avoid modifying summary sheets, skip known ones
    skip_sheets = ['nhắc', 'đóng', 'giảm', 'thiếu', 'bản sao']
    
    for sheet_name in wb.sheetnames:
        if any(x in sheet_name.lower() for x in skip_sheets):
            continue
            
        ws = wb[sheet_name]
        
        # 1. Find the header row and "HỌ VÀ TÊN" column
        header_row = -1
        name_col = -1
        for row in range(1, min(15, ws.max_row + 1)):
            for col in range(1, min(15, ws.max_column + 1)):
                cell_val = str(ws.cell(row=row, column=col).value or '').lower()
                if 'họ và tên' in cell_val or 'tên hs' in cell_val or 'tên học sinh' in cell_val:
                    header_row = row
                    name_col = col
                    break
            if header_row != -1:
                break
                
        if header_row == -1 or name_col == -1:
            print(f"  Skipping sheet '{sheet_name}' (No HỌ VÀ TÊN header found).")
            continue
            
        print(f"  Updating sheet '{sheet_name}'...")
        
        # 2. Insert a new column right after "HỌ VÀ TÊN"
        next_col_val = str(ws.cell(row=header_row, column=name_col + 1).value or '').strip().upper()
        if 'MÃ HS' not in next_col_val:
            ws.insert_cols(name_col + 1)
            ws.cell(row=header_row, column=name_col + 1).value = 'MÃ HS'
            
            # Formatting
            from openpyxl.utils import get_column_letter
            dst_letter = get_column_letter(name_col + 1)
            ws.column_dimensions[dst_letter].width = 12
        else:
            print(f"    MÃ HS column already exists in {sheet_name}.")
            
        # 3. Iterate through rows and populate MÃ HS
        for row in range(header_row + 1, ws.max_row + 1):
            student_name = str(ws.cell(row=row, column=name_col).value or '').strip()
            if not student_name:
                continue
                
            student_name_clean = student_name.lower()
            if student_name_clean in name_to_id:
                ws.cell(row=row, column=name_col + 1).value = name_to_id[student_name_clean]
            else:
                # Try partial match or leave blank
                matched = False
                for k, v in name_to_id.items():
                    if k == student_name_clean or k in student_name_clean or student_name_clean in k:
                        ws.cell(row=row, column=name_col + 1).value = v
                        matched = True
                        break
                if not matched:
                    print(f"    Warning: Could not find ID for '{student_name}'")

    print(f"Saving to {output_path}...")
    wb.save(output_path)
    print(f"Done processing {file_path}.\n")

update_excel_file('C2 ĐIỂM DANH LT1 (26-27).xlsx', 'NEW_C2_DIEM_DANH.xlsx')
update_excel_file('C3 ĐIỂM DANH LT1 (26-27).xlsx', 'NEW_C3_DIEM_DANH.xlsx')
update_excel_file('CẤP 2 HP LT1 26 - 27 .xlsx', 'NEW_C2_HOC_PHI.xlsx')
update_excel_file('CẤP 3 HP LT1 26 - 27 .xlsx', 'NEW_C3_HOC_PHI.xlsx')

print("All files updated successfully.")
